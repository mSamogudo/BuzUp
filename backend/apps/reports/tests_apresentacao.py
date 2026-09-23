# -*- coding: utf-8 -*-
"""O que sai nos documentos, e nao o que esta na base.

Estes testes nasceram de um relatorio que o cliente devolveu com cinco
queixas. Nenhuma era um erro de calculo — eram todas de apresentacao, que e
precisamente o que ninguem testa e toda a gente ve.
"""
from __future__ import annotations

from datetime import datetime, timezone as utc
from decimal import Decimal

from django.test import TestCase
from django.utils import timezone

from apps.reports.exporters import (
    _apresentar, _dinheiro, _e_intocavel, _larguras_das_colunas,
    _rotulo_de_coluna, _rotulo_de_total, _stringify,
)


class ADataEHoraDeMaputo(TestCase):
    """O PDF mostrava a partida das 05:00 como «03:00».

    O Django guarda tudo em UTC e o `strftime` escreve o fuso que a data
    traz. O horario inteiro saia com duas horas de atraso, num documento que
    o operador usa para saber a que horas sai o autocarro.
    """

    def test_uma_data_utc_sai_em_hora_local(self):
        partida = datetime(2026, 9, 25, 3, 0, tzinfo=utc.utc)   # 05:00 em Maputo
        self.assertEqual(_stringify(partida), "25/09/2026 05:00")

    def test_a_da_tarde_tambem(self):
        partida = datetime(2026, 9, 25, 13, 30, tzinfo=utc.utc)  # 15:30 em Maputo
        self.assertEqual(_stringify(partida), "25/09/2026 15:30")

    def test_uma_data_sem_fuso_nao_rebenta(self):
        self.assertEqual(_stringify(datetime(2026, 9, 25, 5, 0)), "25/09/2026 05:00")

    def test_o_dia_vem_a_frente_do_mes(self):
        """Em Mocambique le-se 25/09. «09/25» num relatorio de fronteira e um
        engano a espera de acontecer."""
        self.assertTrue(_stringify(datetime(2026, 9, 25, 5, 0)).startswith("25/09"))


class ODinheiro(TestCase):
    def test_le_se_como_em_mocambique(self):
        self.assertEqual(_dinheiro("68650.00"), "68 650,00 MZN")
        self.assertEqual(_dinheiro("1650"), "1 650,00 MZN")
        self.assertEqual(_dinheiro(Decimal("1234567.5")), "1 234 567,50 MZN")

    def test_o_separador_e_um_espaco_que_o_pdf_sabe_desenhar(self):
        """O espaco fino inquebravel (U+202F) nao existe no Helvetica: saiu um
        quadrado preto no meio do total — «68<quadrado>650,00»."""
        self.assertNotIn(" ", _dinheiro("68650.00"))
        self.assertNotIn(" ", _dinheiro("68650.00"))

    def test_o_que_nao_e_numero_passa_intacto(self):
        self.assertEqual(_dinheiro("—"), "—")

    def test_na_tabela_vai_sem_a_moeda(self):
        """A moeda esta no cabecalho da coluna. Repetida em 62 linhas so
        rouba largura ao numero."""
        self.assertEqual(_apresentar("amount", "1650.00"), "1 650,00")
        self.assertEqual(_rotulo_de_coluna("amount", "Valor"), "Valor (MZN)")

    def test_no_excel_vai_como_NUMERO(self):
        """Formatar o dinheiro em texto deixava a folha bonita e impossivel
        de somar — e somar a coluna e a primeira coisa que quem a recebe faz."""
        valor = _apresentar("amount", "1650.00", para="xlsx")
        self.assertIsInstance(valor, float)
        self.assertEqual(valor, 1650.0)


class OsEstados(TestCase):
    def test_saem_por_escrito(self):
        self.assertEqual(_apresentar("status", "confirmed"), "Confirmado")
        self.assertEqual(_apresentar("status", "failed"), "Falhado")

    def test_um_estado_desconhecido_ainda_se_le(self):
        self.assertEqual(_apresentar("status", "meio_pago"), "Meio pago")

    def test_vazio_nao_se_confunde_com_zero(self):
        """Uma celula em branco le-se como «nao houve». Aqui quer dizer «nao
        soubemos dizer», e o documento tem de o distinguir."""
        self.assertEqual(_apresentar("status", ""), "—")


class OsRotulosDosCartoes(TestCase):
    def test_estao_em_portugues(self):
        """Sairam «COUNT» e «CONFIRMED COUNT» — nomes das chaves do
        dicionario, no primeiro sitio que se le no documento."""
        self.assertEqual(_rotulo_de_total("count"), "Linhas")
        self.assertEqual(_rotulo_de_total("confirmed_count"), "Confirmadas")
        self.assertEqual(_rotulo_de_total("total_amount"), "Total")

    def test_uma_chave_nova_nao_sai_com_underscores(self):
        self.assertEqual(_rotulo_de_total("algo_novo"), "Algo novo")


class AsLarguras(TestCase):
    """«1650…» num relatorio financeiro nao e um numero mais curto — e outro
    numero."""

    class TelaFalsa:
        @staticmethod
        def stringWidth(txt, font, size):
            return len(txt or "") * size * 0.5

    def test_o_dinheiro_e_as_datas_sao_intocaveis(self):
        for key in ("amount", "fare_amount", "created_at", "ida_at", "status"):
            self.assertTrue(_e_intocavel(key), key)
        for key in ("reference", "route_code", "passageiro"):
            self.assertFalse(_e_intocavel(key), key)

    def test_a_coluna_do_dinheiro_recebe_o_que_precisa(self):
        """Mesmo ao lado de referencias longas, que antes lhe comiam a
        largura toda."""
        colunas = [
            ("reference", "Pagamento"), ("sale_reference", "Venda"),
            ("route_code", "Rota"), ("amount", "Valor"),
        ]
        linhas = [{
            "reference": "PAY-AS-57735E7CA2914B2C3D4E5F6A7B8C9D",
            "sale_reference": "AS-57735E7CA2914B2C3D4E5F6A7B8C9D",
            "route_code": "RT-MAPUTO-X-NELSPRUIT",
            "amount": "1650.00",
        }]
        c = self.TelaFalsa()
        larguras = _larguras_das_colunas(c, colunas, linhas, 800.0)

        i = [k for k, _ in colunas].index("amount")
        texto = _apresentar("amount", "1650.00")
        precisa = c.stringWidth(texto, "Helvetica", 8)
        self.assertGreaterEqual(larguras[i], precisa,
                                "a coluna do dinheiro tem de caber sem cortar")

    def test_a_soma_das_larguras_cabe_na_pagina(self):
        colunas = [(f"c{i}", f"Coluna {i}") for i in range(8)] + [("amount", "Valor")]
        linhas = [{f"c{i}": "x" * 20 for i in range(8)} | {"amount": "1650.00"}]
        larguras = _larguras_das_colunas(self.TelaFalsa(), colunas, linhas, 800.0)
        self.assertLessEqual(sum(larguras), 800.0 + 0.5)


class OQueCadaRelatorioConta(TestCase):
    """Cada relatorio diz o que conta e o que deixa de fora.

    O cliente comparou o total de um relatorio com o cartao «Receita de
    transporte» do painel e encontrou 10 900,00 MZN de diferenca. Nao era
    erro: o painel soma bilhetes MAIS validacoes de cartao, e o relatorio so
    cobre um dos dois.

    Um numero que nao diz o que conta obriga quem o le a desconfiar de todos
    os outros.
    """

    def test_todos_os_relatorios_declaram_o_escopo(self):
        from apps.reports.builder import REGISTRY

        for chave, spec in REGISTRY.items():
            self.assertTrue(getattr(spec, "escopo", ""),
                            f"o relatorio «{chave}» nao diz o que conta")

    def test_o_escopo_diz_o_que_NAO_conta(self):
        """A metade que importa. «Conta bilhetes» nao ajuda ninguem a
        perceber porque e que o painel mostra mais."""
        from apps.reports.builder import REGISTRY

        for chave in ("sales", "tickets", "topups", "validations"):
            self.assertIn("Nao", REGISTRY[chave].escopo,
                          f"o escopo de «{chave}» nao diz o que fica de fora")

    def test_o_escopo_chega_ao_documento(self):
        from apps.reports.builder import REGISTRY
        from apps.reports.exporters import render_xlsx
        import io as _io
        from openpyxl import load_workbook

        spec = REGISTRY["sales"]
        dados = render_xlsx(
            title=spec.title, period_from="01/09/2026", period_to="30/09/2026",
            columns=spec.columns, rows=[], totals={}, escopo=spec.escopo,
        )
        wb = load_workbook(_io.BytesIO(dados))
        ws = wb["Resumo"]
        self.assertEqual(ws["A5"].value, spec.escopo)
        self.assertTrue(str(ws["A4"].value or "").startswith("Gerado em"),
                        "o «gerado em» nao pode ser apagado pelo escopo")


class AEscolhaDeColunas(TestCase):
    def setUp(self):
        from apps.reports.api.views import _colunas_pedidas
        from apps.reports.builder import REGISTRY

        self.escolher = _colunas_pedidas
        self.spec = REGISTRY["sales"]

    def test_sem_pedido_vao_todas(self):
        """Quem ja tem um link guardado recebe o mesmo de sempre."""
        self.assertEqual(self.escolher(self.spec, None), self.spec.columns)
        self.assertEqual(self.escolher(self.spec, ""), self.spec.columns)

    def test_a_ordem_e_a_do_relatorio_e_nao_a_do_pedido(self):
        """Se fosse a do pedido, duas pessoas com as mesmas colunas recebiam
        documentos diferentes e ninguem os podia comparar."""
        escolhidas = self.escolher(self.spec, "status,created_at,amount")
        chaves = [k for k, _ in escolhidas]
        ordem_original = [k for k, _ in self.spec.columns if k in set(chaves)]
        self.assertEqual(chaves, ordem_original)

    def test_uma_coluna_que_ja_nao_existe_ignora_se(self):
        """O pedido vem de um URL que alguem guardou nos favoritos."""
        escolhidas = self.escolher(self.spec, "created_at,coluna_que_morreu")
        self.assertEqual([k for k, _ in escolhidas], ["created_at"])

    def test_nenhuma_valida_devolve_tudo(self):
        """Um documento vazio nao se distingue de «nao houve movimento»."""
        self.assertEqual(self.escolher(self.spec, "nada,disto,existe"), self.spec.columns)
